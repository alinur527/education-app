"""Extract factual official section/topic titles, preserving source page/codes.

Not a content importer. Explicit boundaries handle PDF table text-box artifacts;
they come from inspected PDF coordinates, not from inferred subject knowledge.
"""
import hashlib
import html
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote

import openpyxl
import pdfplumber

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / 'content/inputs'
DOWNLOADS = ROOT / 'test-results/source-research'
OUTPUT = ROOT / 'content/research'
BOUNDARIES = {
    'SPEC-05': [51.3, 72.6, 187.4, 214.3, 547.1],
    'SPEC-07': [56.7, 79.8, 175.1, 200.9, 538.6],
    'SPEC-08': [56.7, 70.7, 176.9, 206.6, 538.6],
    'SPEC-11': [56.6, 79.7, 209.7, 232.3, 547.3],
    'SPEC-16': [56.6, 77.5, 217.8, 245.7, 558.3],
    'SPEC-17': [56.6, 77.5, 209.7, 241.3, 560.7],
    'SPEC-20': [55.9, 83.9, 170.0, 193.4, 551.7],
    'SPEC-21': [56.6, 78.1, 155.9, 177.1, 542.7],
    'SPEC-29': [56.6, 83.4, 210.9, 232.3, 543.7],
}


def clean(value):
    return re.sub(r'\s+', ' ', value or '').strip()


def code(value):
    value = re.sub(r'\s+', '', value or '')
    return value if re.fullmatch(r'\d{1,2}', value) else None


def world_history_kk_rows(page, page_number):
    """The KK world-history PDF omits left borders on continued pages.

    Read its four visually inspected column areas rather than dropping those
    pages because pdfplumber detects only a three-column table there.
    """
    if page_number > 5:
        return []
    bounds = BOUNDARIES['SPEC-20']
    top = 263 if page_number == 1 else 56.7
    bottom = [749, 781.7, 737.5, 783.4, 720.8][page_number - 1]
    def markers(left, right):
        chars = [ch for ch in page.chars if left < ch['x0'] < right and top < ch['top'] < bottom and ch['text'].strip()]
        lines = {}
        for ch in chars:
            y = next((old for old in lines if abs(old - ch['top']) < 2), ch['top'])
            lines.setdefault(y, []).append(ch)
        found = []
        for y, line in sorted(lines.items()):
            value = ''.join(ch['text'] for ch in sorted(line, key=lambda ch: ch['x0']))
            if code(value):
                found.append((y, code(value)))
        return found
    sections = markers(bounds[0], bounds[1])
    topics = markers(bounds[2], bounds[3])
    rows = []
    if topics and topics[0][0] > top + 12:
        continued = clean(page.crop((bounds[3], top, bounds[4], topics[0][0] - 1)).extract_text())
        if continued:
            rows.append(['', '', '', continued])
    for idx, (y, tc) in enumerate(topics):
        end = topics[idx + 1][0] - 1 if idx + 1 < len(topics) else bottom
        topic = clean(page.crop((bounds[3], y - 1, bounds[4], end)).extract_text())
        found = next(((j, sc) for j, (sy, sc) in enumerate(sections) if abs(sy - y) < 3), None)
        if found:
            j, sc = found
            section_end = sections[j + 1][0] - 1 if j + 1 < len(sections) else bottom
            title = clean(page.crop((bounds[1], y - 1, bounds[2], section_end)).extract_text())
            rows.append([sc, title, tc, topic])
        else:
            rows.append(['', '', tc, topic])
    return rows


def extract(spec):
    sid = spec['id']
    manifest = json.loads((DOWNLOADS / f'{sid}.manifest.json').read_text(encoding='utf-8'))
    sections = []
    problems = []
    current = None
    last_topic = None
    pages_text = []
    with pdfplumber.open(DOWNLOADS / f'{sid}.pdf') as pdf:
        for index, page in enumerate(pdf.pages):
            text = page.extract_text() or ''
            pages_text.append({'page': index + 1, 'text': text})
            settings = {'vertical_strategy': 'explicit', 'explicit_vertical_lines': BOUNDARIES[sid], 'horizontal_strategy': 'lines'} if sid in BOUNDARIES else {}
            tables = [world_history_kk_rows(page, index + 1)] if sid == 'SPEC-20' else page.extract_tables(settings)
            for table in tables:
                if not table:
                    continue
                if len(table[0]) != 4:
                    continue
                for row in table:
                    section_code, section_title, topic_code, topic_title = row
                    if clean(section_code) == '№' or clean(topic_code) == '№':
                        continue
                    sc, tc = code(section_code), code(topic_code)
                    if sid == 'SPEC-21' and index == 4 and tc == '38' and not sc:
                        current = {'officialCode': None, 'editorialKey': 'uncoded-section-page-5-topic-38', 'title': clean(section_title), 'pages': [5], 'topics': [], 'sourceAnomaly': 'The section number cell is visibly blank in the official Russian PDF; do not invent code 12.'}
                        sections.append(current)
                    if sc and clean(section_title):
                        current = {'officialCode': sc, 'title': clean(section_title), 'pages': [index + 1], 'topics': []}
                        sections.append(current)
                    if tc and clean(topic_title):
                        if current is None:
                            problems.append(f'page {index + 1}: topic {tc} has no identified section')
                            continue
                        last_topic = {'officialCode': tc, 'title': clean(topic_title), 'pages': [index + 1], 'sourceId': sid, 'extractionStatus': 'EXTRACTED_FROM_OFFICIAL_TABLE'}
                        current['topics'].append(last_topic)
                        if index + 1 not in current['pages']:
                            current['pages'].append(index + 1)
                    elif not sc and not tc and clean(topic_title) and last_topic and not clean(topic_code):
                        # A broad topic may continue over the page break.
                        if clean(topic_title) not in ('Тема', 'Тақырып', 'Тақырыптар', 'Topic'):
                            last_topic['title'] += ' ' + clean(topic_title)
                            if index + 1 not in last_topic['pages']:
                                last_topic['pages'].append(index + 1)
                    elif tc and not clean(topic_title):
                        problems.append(f'page {index + 1}: topic {tc} empty after extraction')
    topics = [topic for section in sections for topic in section['topics']]
    if sid == 'SPEC-30':
        problems.append('Kazakh literature topic 23 is visibly clipped in PDF page 2 (the title stops after «Біз»); do not silently reconstruct the missing ending. Topic 24 was visually transcribed from the rendered page because its text layer overlaps topic 23.')
        for topic in topics:
            if topic['officialCode'] == '23':
                topic['extractionStatus'] = 'UNCERTAIN_SOURCE_TITLE_VISIBLY_CLIPPED'
            if topic['officialCode'] == '24':
                topic['extractedTitleBeforeReview'] = topic['title']
                topic['title'] = 'Т. Нұрмағамбетов "Анасын сағынған бала", Н. Ақыш "Нағыз әже қайда",'
                topic['extractionStatus'] = 'VISUALLY_TRANSCRIBED_FROM_OFFICIAL_PAGE'
                topic['visualEvidence'] = {'page': 2, 'renderPath': 'test-results/source-research/spec30-page2.png', 'reviewKind': 'MODEL_VISUAL_READING_NOT_HUMAN_REVIEW'}
    codes = [int(topic['officialCode']) for topic in topics]
    if codes != list(range(1, max(codes, default=0) + 1)):
        problems.append('Topic codes are not a complete increasing 1..N sequence; review required')
    full = '\n'.join(page['text'] for page in pages_text)
    first = pages_text[0]['text']
    evidence = next((line for line in first.splitlines() if '2026' in line), None)
    applies = 2026 if evidence and ('использования' in evidence.lower() or 'қолдану' in evidence.lower()) else None
    item = {
        'sourceId': sid, 'subjectRu': spec['subject'], 'category': 'MANDATORY' if spec['subject'] in ('Математическая грамотность', 'Грамотность чтения', 'История Казахстана') else 'PROFILE',
        'documentLanguage': spec['language'], 'programmeVariant': spec['variant'],
        'examVersion': 'NTC_FROM_2026' if applies else None, 'appliesFrom': applies,
        'appliesFromEvidence': {'page': 1, 'text': evidence},
        'url': manifest['url'], 'resolvedUrl': manifest['resolvedUrl'], 'sha256': manifest['sha256'],
        'pageCount': manifest['pageCount'], 'sectionCount': len(sections), 'topicCount': len(topics),
        'extractionStatus': 'REQUIRES_REVIEW' if problems else 'COMPLETE_CODE_SEQUENCE_EXTRACTED',
        'extractionWarnings': problems, 'sections': sections,
        'requirementsEvidence': [],
    }
    for page in pages_text:
        for line in page['text'].splitlines():
            if any(key in line.lower() for key in ('40 тест', '20 тест', '10 тест', '25 тест', '5 тест', 'минут', 'тест тапсырмас', 'минут', 'контекст', 'сәйкесті', 'соответств', '50%', '30%', '20%')):
                item['requirementsEvidence'].append({'page': page['page'], 'text': line})
    return item


def consistency(catalog):
    workbook = openpyxl.load_workbook(INPUT / 'ENT_content_catalog_2026-10-04.xlsx', read_only=True, data_only=True)
    xlsx_text = '\n'.join(str(value) for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row if value is not None)
    markdown = (INPUT / 'ENT_research_NotebookLM_2026-10-04.md').read_text(encoding='utf-8-sig')
    normalize = lambda value: unquote(html.unescape(value))
    normalized_xlsx, normalized_md = normalize(xlsx_text), normalize(markdown)
    mismatches = []
    for key in ('specifications', 'officialSamples', 'recommendedBookLists', 'officialReferences', 'webSources'):
        for item in catalog[key]:
            for origin, body in [('XLSX', normalized_xlsx), ('MARKDOWN', normalized_md)]:
                if normalize(item['url']) not in body:
                    mismatches.append({'sourceId': item['id'], 'origin': origin, 'issue': 'JSON URL absent from this representation'})
    return {'inputSha256': {file.name: hashlib.sha256(file.read_bytes()).hexdigest() for file in sorted(INPUT.iterdir()) if file.is_file()}, 'sheets': workbook.sheetnames, 'urlMismatches': mismatches}


def main():
    catalog = json.loads((INPUT / 'ENT_sources_2026-10-04.json').read_text(encoding='utf-8-sig'))
    documents = [extract(spec) for spec in catalog['specifications']]
    result = {'schema': 'education-app-official-curriculum-v1', 'checkedAt': datetime.now(timezone.utc).isoformat(), 'scope': 'NTC full-duration admission; official source document variants remain separate', 'editorialApproval': False, 'programmes': documents, 'inputConsistency': consistency(catalog)}
    (OUTPUT / 'official-curriculum.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    for item in documents:
        print(item['sourceId'], item['sectionCount'], item['topicCount'], item['extractionWarnings'])
    print('CONSISTENCY', result['inputConsistency'])


if __name__ == '__main__':
    main()
